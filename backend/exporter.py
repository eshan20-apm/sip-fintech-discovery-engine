import io
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.database import get_db
from backend.analytics.benchmark_eval import get_evaluation_metrics

def format_worksheet(worksheet, header_color="1E293B", accent_color="38BDF8"):
    """Applies professional formatting, column auto-width, freeze panes, and clean headers."""
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Style header row
    for col_idx in range(1, worksheet.max_column + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Freeze header
    worksheet.freeze_panes = "A2"

    # Style data rows and auto-fit widths
    for col in worksheet.columns:
        col_letter = col[0].column_letter
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
            if cell.row > 1:
                cell.font = Font(name="Calibri", size=10)
                # Align text columns vs IDs/dates
                if col_letter in ['A', 'B', 'C', 'D', 'K']:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        worksheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 55)

def export_posts_to_excel(posts: List[Dict[str, Any]]) -> bytes:
    """
    Generates a 3-tab executive research workbook using pure openpyxl (zero heavy dependencies):
    1. 'Active Research Feed': Filtered comments matching the prompt/platform view
    2. 'Golden Benchmark Set': Hand-verified truth set with PM rationales
    3. 'Reliability & Kappa Metrics': Statistical evaluation (Cohen's Kappa κ, accuracy rates)
    """
    wb = Workbook()

    # ---------------- TAB 1: ACTIVE RESEARCH FEED ----------------
    ws1 = wb.active
    ws1.title = "Active Research Feed"

    headers1 = [
        "Post ID", "Platform", "Source", "Published Date",
        "Dynamic Theme (2-3 Words)", "Journey Stage", "Action Taken",
        "Investor Persona", "Responsible Entity", "Verbatim Evidence Quote",
        "Status", "RRF Relevance Score", "PM Notes", "Raw Feedback", "Original Source URL"
    ]
    ws1.append(headers1)

    for p in posts:
        is_golden = p.get("is_hand_labelled") == 1
        theme = p.get("golden_theme") if is_golden else p.get("ai_theme")
        journey_stage = p.get("golden_journey_stage") if is_golden else p.get("ai_journey_stage")
        what_they_did = p.get("golden_what_they_did") if is_golden else p.get("ai_what_they_did")
        who_they_are = p.get("golden_who_they_are") if is_golden else p.get("ai_who_they_are")
        who_can_fix_it = p.get("golden_who_can_fix_it") if is_golden else p.get("ai_who_can_fix_it")
        evidence_quote = p.get("golden_evidence_quote") if (is_golden and p.get("golden_evidence_quote")) else p.get("ai_evidence_quote")

        ws1.append([
            p.get("id"),
            p.get("platform"),
            p.get("source"),
            p.get("published_at", "")[:10] if p.get("published_at") else "Live",
            theme or "Unclear",
            journey_stage or "Active SIP Phase",
            what_they_did or "Unsure",
            who_they_are or "Retail investor",
            who_can_fix_it or "Fintech Platform",
            evidence_quote or "",
            "Golden Benchmark" if is_golden else "AI Tagged",
            p.get("rrf_score", "N/A"),
            p.get("golden_notes", ""),
            p.get("raw_text", ""),
            p.get("url", "")
        ])

    format_worksheet(ws1, header_color="1E293B")

    # ---------------- TAB 2: GOLDEN BENCHMARK SET ----------------
    ws2 = wb.create_sheet(title="Golden Benchmark Set")
    headers2 = [
        "Post ID", "Platform", "Source", "Date",
        "Verified Theme", "Verified Journey Stage", "Verified Action Taken",
        "Investor Persona", "Responsible Entity", "Verbatim Evidence Quote",
        "PM Benchmark Rationale / Notes", "Raw Comment Text", "Source Link"
    ]
    ws2.append(headers2)

    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("""
    SELECT 
        p.id, p.platform, p.source, p.published_at, p.raw_text, p.url,
        g.theme, g.journey_stage, g.what_they_did, g.who_they_are, g.who_can_fix_it,
        g.evidence_quote, g.notes, g.updated_at
    FROM golden_tags g
    JOIN posts p ON g.post_id = p.id
    ORDER BY g.updated_at DESC
    """)
    rows_golden = [dict(r) for r in cursor.fetchall()]
    conn.close()

    for g in rows_golden:
        ws2.append([
            g.get("id"),
            g.get("platform"),
            g.get("source"),
            g.get("published_at", "")[:10] if g.get("published_at") else "",
            g.get("theme"),
            g.get("journey_stage"),
            g.get("what_they_did"),
            g.get("who_they_are"),
            g.get("who_can_fix_it"),
            g.get("evidence_quote") or "",
            g.get("notes") or "",
            g.get("raw_text"),
            g.get("url")
        ])

    format_worksheet(ws2, header_color="065F46")

    # ---------------- TAB 3: RELIABILITY & KAPPA METRICS ----------------
    ws3 = wb.create_sheet(title="Reliability & Kappa Metrics")
    headers3 = ["Metric", "Value", "Unit / Scale"]
    ws3.append(headers3)

    eval_metrics = get_evaluation_metrics()
    metrics_summary = [
        ["Overall Cohen's Kappa (κ)", str(eval_metrics.get("overall_kappa")), "-1.0 to +1.0 (Chance-Adjusted)"],
        ["Reliability Interpretation", str(eval_metrics.get("interpretation")), "Landis & Koch (1977) Standard"],
        ["Overall Observed Raw Agreement (Pₒ)", f"{eval_metrics.get('overall_agreement_rate')}%", "Exact label match %"],
        ["Total Golden Benchmark Posts", str(eval_metrics.get("total_benchmarked")), "Hand-verified records"],
        ["Theme Classification Kappa", str(eval_metrics.get("fields", {}).get("theme", {}).get("kappa")), "Field Reliability"],
        ["Theme Raw Accuracy", f"{eval_metrics.get('fields', {}).get('theme', {}).get('accuracy')}%", "Exact match %"],
        ["Journey Stage Kappa", str(eval_metrics.get("fields", {}).get("journey_stage", {}).get("kappa")), "Field Reliability"],
        ["Journey Stage Raw Accuracy", f"{eval_metrics.get('fields', {}).get('journey_stage', {}).get('accuracy')}%", "Exact match %"],
        ["Action Taken Kappa", str(eval_metrics.get("fields", {}).get("what_they_did", {}).get("kappa")), "Field Reliability"],
        ["Action Taken Raw Accuracy", f"{eval_metrics.get('fields', {}).get('what_they_did', {}).get('accuracy')}%", "Exact match %"],
        ["Who Can Fix It Kappa", str(eval_metrics.get("fields", {}).get("who_can_fix_it", {}).get("kappa")), "Field Reliability"],
        ["Who Can Fix It Raw Accuracy", f"{eval_metrics.get('fields', {}).get('who_can_fix_it', {}).get('accuracy')}%", "Exact match %"]
    ]

    for m in metrics_summary:
        ws3.append(m)

    format_worksheet(ws3, header_color="312E81")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
